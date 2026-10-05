from __future__ import annotations

import json
import re
import shutil
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path.cwd()
BATCH = ROOT / ".firecrawl/linkedin-company-batch-2026-08-13"
OUTPUT = ROOT / "临时/批量客户调查-2026-08-13"
PROFILE_ROOT = OUTPUT / "档案"
SUMMARY_MD = OUTPUT / "全部客户取舍原因.md"
SUMMARY_XLSX = ROOT / "临时/_批量汇总_2026-08-13.xlsx"


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def unwrap_agent_output(raw: Any) -> list[dict[str, Any]]:
    current = raw
    for _ in range(5):
        if isinstance(current, str):
            try:
                current = json.loads(current)
                continue
            except json.JSONDecodeError:
                return []
        if isinstance(current, dict) and isinstance(current.get("companies"), list):
            return current["companies"]
        if isinstance(current, dict) and "data" in current:
            current = current["data"]
            continue
        break
    return []


def md(value: Any) -> str:
    if value is None or value == "":
        return "待确认"
    return str(value).replace("|", "\\|").replace("\n", " ").strip()


def bullets(items: Any, empty: str = "- 待确认") -> str:
    if not items:
        return empty
    lines = []
    for item in items:
        if isinstance(item, dict):
            text = "；".join(f"{key}: {value}" for key, value in item.items() if value not in (None, "", []))
        else:
            text = str(item)
        lines.append(f"- {text}")
    return "\n".join(lines)


def safe_filename(value: str) -> str:
    cleaned = re.sub(r"[\\/:*?\"<>|]+", "-", value).strip(" .-")
    return re.sub(r"\s+", "-", cleaned)[:120] or "company"


def country_folder(company: dict[str, Any], source: dict[str, Any]) -> str:
    value = f"{company.get('country', '')} {company.get('city_region', '')} {source.get('location', '')}".lower()
    if any(token in value for token in ["canada", "ontario", "quebec", "british columbia", "alberta", "saskatchewan", "manitoba", "nova scotia", "winnipeg", "vancouver", "toronto", "ottawa", "montreal", "mississauga", "calgary"]):
        return "加拿大"
    if any(token in value for token in ["australia", "sydney", "melbourne", " nsw", "victoria"]):
        return "澳大利亚"
    if any(token in value for token in ["england", "united kingdom", " uk", "scotland", "wales", "northern ireland", "surrey", "essex", "london", "manchester", "leeds", "hampshire", "kent", "cheshire", "yorkshire", "midlands"]):
        return "英国"
    if any(token in value for token in ["ireland", "cork"]):
        return "爱尔兰"
    if any(token in value for token in ["netherlands", "holland", "amsterdam", "delft"]):
        return "荷兰"
    if any(token in value for token in ["germany", "dusseldorf"]):
        return "德国"
    if any(token in value for token in ["poland", "warsaw"]):
        return "波兰"
    if any(token in value for token in ["malta", "attard"]):
        return "马耳他"
    if any(token in value for token in ["united states", "usa", "new york", "delaware", "ohio", "las vegas", "california"]):
        return "美国"
    return "其他"


def score_table(title: str, total: Any, breakdown: dict[str, Any], rows: list[tuple[str, str, int]]) -> str:
    lines = [f"### {title}: {total}/100", "", "| 维度 | 满分 | 得分 |", "|---|---:|---:|"]
    for label, key, maximum in rows:
        lines.append(f"| {label} | {maximum} | {breakdown.get(key, '待确认')} |")
    return "\n".join(lines)


def make_profile(company: dict[str, Any], source: dict[str, Any]) -> str:
    contacts = company.get("contacts") or []
    primary = contacts[0] if contacts else {}
    needs = company.get("potential_procurement_needs") or []
    pains = company.get("pain_points") or []
    angles = sorted(company.get("china_supplier_entry_angles") or [], key=lambda item: item.get("rank", 99))
    products = company.get("product_lines") or []
    sources = company.get("research_sources") or []
    customs = company.get("customs_platform_checks") or {}

    product_rows = [
        f"| {md(item.get('category'))} | {md(item.get('materials'))} | {md(item.get('procurement_model'))} | {md(item.get('metal_presence'))} | {md(item.get('evidence'))} |"
        for item in products
    ] or ["| 待确认 | 待确认 | 待确认 | Unknown | 证据不足 |"]
    contact_rows = [
        f"| {index} | {md(item.get('name'))} | {md(item.get('title'))} | {md(item.get('decision_power'))} | {md(item.get('evidence_strength'))} | {md(item.get('contact_channel'))} | {md(item.get('verification'))} |"
        for index, item in enumerate(contacts, 1)
    ] or ["| 1 | 待确认 | 待确认 | 待确认 | Weak | 官网表单/通用渠道 | unverified |"]
    pain_rows = [
        f"| {md(item.get('pain'))} | {md(item.get('evidence_or_inference'))} | {md(item.get('confidence'))} |"
        for item in pains
    ] or ["| 待确认 | 当前证据不足 | Low |"]
    need_rows = [
        f"| {md(item.get('need'))} | {md(item.get('probability'))} | {md(item.get('evidence'))} |"
        for item in needs
    ] or ["| 待确认 | Unknown | 当前证据不足 |"]
    angle_sections = []
    for item in angles[:3]:
        angle_sections.append(
            f"### Angle {item.get('rank', len(angle_sections) + 1)}：{md(item.get('angle'))}\n\n"
            f"- **为何成立**：{md(item.get('why'))}\n"
            f"- **产品/场景配对**：{md(item.get('product_scene'))}\n"
            f"- **风险**：{md(item.get('risk'))}"
        )
    if not angle_sections:
        angle_sections.append("### Angle 1：待确认\n\n当前证据不足，不应主动承诺产品方案。")

    primary_name = primary.get("name") or ""
    first_name = "there" if not primary_name or "通用联系" in primary_name else primary_name.split()[0]
    invite = (
        f"Hi {first_name}, I work with a China-based custom metal gift maker. "
        f"I noticed {company.get('company')} works in branded merchandise. "
        "We support planned metal gift projects beyond standard catalogue items, MOQ 100. Open to connecting?"
    )
    invite = invite[:297] + "..." if len(invite) > 300 else invite
    best_angle = angles[0].get("angle") if angles else "planned custom metal gift projects"
    website = company.get("official_website") or source.get("official_url") or "待确认"
    location = company.get("city_region") or source.get("location") or "待确认"
    role = company.get("supply_chain_role") or "Unknown"

    fit_rows = [
        ("产品匹配", "product_fit", 20), ("买家角色匹配", "buyer_role_fit", 15),
        ("证据强度", "evidence_strength", 15), ("采购潜力", "procurement_potential", 15),
        ("可触达性", "reachability", 10), ("市场价值", "market_value", 10),
        ("供应商匹配", "supplier_fit", 10), ("风险调整", "risk_adjustment", -15),
    ]
    value_rows = [
        ("年采购潜力", "annual_procurement", 20), ("利润质量", "margin_quality", 15),
        ("战略市场价值", "strategic_value", 15), ("合作深度", "cooperation_depth", 15),
        ("转化可及性", "conversion_access", 10), ("供应商能力匹配", "capability_fit", 10),
        ("关系效率", "relationship_efficiency", 10), ("风险调整", "risk_adjustment", -15),
    ]

    return f"""# {company.get('company')} — 采购背景档案

> 调查日期：2026-08-13  
> Gate 2：{company.get('gate2_grade', '待确认')}  
> 目标客户：{'是' if company.get('target_customer') else '否'}  
> 所有推断均与事实证据分开；未知项保留为待确认。

**网站**：{website}  
**位置**：{location}  
**首选联系人**：{md(primary.get('name'))}，{md(primary.get('title'))}  
**邮箱**：{md(primary.get('email'))}（{md(primary.get('verification'))}）  
**电话**：{md(primary.get('phone'))}

## 采购背景档案（速览）

- **核心业务方向**：{md(company.get('core_business'))}
- **主要产品线**：{md('、'.join(item.get('category', '') for item in products if item.get('category')))}
- **潜在采购需求**：{md('；'.join(f"{item.get('need')}（{item.get('probability')}）" for item in needs))}
- **供应链可能痛点**：{md('；'.join(item.get('pain', '') for item in pains if item.get('pain')))}
- **中国供应商最佳切入角度**：{md(best_angle)}

## 是否值得开发

**结论**：{md(company.get('decision_reason'))}

- Fit：{company.get('fit_score', '待确认')}/100
- Value：{company.get('value_score', '待确认')}/100
- 价值梯队：{md(company.get('value_tier'))}
- 账户类型：{md(company.get('account_type'))}
- 建议动作：{md(company.get('recommended_action'))}

## 供应链位置

```text
Manufacturer → Exporter → Importer → Distributor → Dealer/Decorator → End buyer → Recipient
                                     ↑ {company.get('company')}：{role}
```

- **角色判断**：{md(company.get('supply_chain_position_explanation'))}
- **可能向谁采购**：{md(company.get('buys_from'))}
- **主要卖给谁**：{md(company.get('sells_to'))}
- **商业模式**：{md('、'.join(company.get('business_model_tags') or []))}

## 公司与业务证据

| 字段 | 调查结果 |
|---|---|
| 成立年份 | {md(company.get('founded_year'))} |
| 员工规模 | {md(company.get('employee_range'))} |
| 终端客户行业 | {md('、'.join(company.get('client_industries') or []))} |
| 客户/案例 | {md('；'.join(company.get('named_clients_or_cases') or []))} |
| 协会/展会 | {md('；'.join(company.get('association_and_trade_show_evidence') or []))} |
| 采购/进口信号 | {md('；'.join(company.get('sourcing_import_signals') or []))} |

## 产品品类地图

| 品类 | 材质 | 采购方式 | 金属存在 | 证据 |
|---|---|---|---|---|
{chr(10).join(product_rows)}

### 产品四问

| 问题 | 调查结果 |
|---|---|
| 是否做金属产品 | {md('是/部分' if company.get('metal_products') else '未确认')} |
| 具体金属产品 | {md('、'.join(company.get('metal_products') or []))} |
| 是否已有企业礼品套件 | {md(company.get('executive_gift_sets'))} |
| 是否已有企业礼品单品 | {md(company.get('executive_gift_singles'))} |

## 潜在采购需求

| 需求 | 可能性 | 依据 |
|---|---|---|
{chr(10).join(need_rows)}

## 供应链痛点

| 痛点 | 证据或推断 | 置信度 |
|---|---|---|
{chr(10).join(pain_rows)}

## 联系人优先级矩阵

| Priority | 联系人 | 职位 | 决策力 | 证据强度 | 渠道 | 验证状态 |
|---:|---|---|---|---|---|---|
{chr(10).join(contact_rows)}

## 海关与 B2B 足迹

| 平台 | 结果 |
|---|---|
{chr(10).join(f'| {md(key)} | {md(value)} |' for key, value in customs.items()) if customs else '| 全部平台 | 未完成或无法访问 |'}

## 评分

{score_table('Fit Score', company.get('fit_score', '待确认'), company.get('fit_breakdown') or {}, fit_rows)}

{score_table('Value Score', company.get('value_score', '待确认'), company.get('value_breakdown') or {}, value_rows)}

## 三个切入角度

{chr(10).join(angle_sections)}

**供应商切换假设**：{md(company.get('supplier_switch_hypothesis'))}

## 结构性风险

{bullets(company.get('structural_risks'))}

## LinkedIn 外展定位

> {invite}

字符数：{len(invite)}

首封邮件建议围绕 **{md(best_angle)}** 展开；对本地 dealer/reseller 明确定位为 planned custom metal projects beyond standard catalogue items，不要求替换其日常目录供应链。

## 未确认事项

{bullets(company.get('unknowns'))}

## 研究来源

{chr(10).join(f"- [{md(item.get('title'))}]({item.get('url')}) — {md(item.get('supports'))}" for item in sources) if sources else '- 暂无可验证来源，需复核'}
"""


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    if PROFILE_ROOT.exists():
        shutil.rmtree(PROFILE_ROOT)
    PROFILE_ROOT.mkdir(parents=True, exist_ok=True)
    audited = BATCH / "gate1-results-audited.json"
    gate1 = load_json(audited if audited.exists() else BATCH / "gate1-results.json")
    source_by_id = {int(item["id"]): item for item in gate1}

    researched: dict[int, dict[str, Any]] = {}
    invalid_deep_files: list[str] = []
    local_analysis = BATCH / "local-analysis/companies.json"
    if local_analysis.exists():
        try:
            for item in load_json(local_analysis).get("companies", []):
                researched[int(item["id"])] = item
        except Exception:
            invalid_deep_files.append(str(local_analysis.relative_to(BATCH)))
    for path in sorted((BATCH / "deep").glob("batch-*.json")):
        try:
            companies = unwrap_agent_output(load_json(path))
        except Exception:
            invalid_deep_files.append(path.name)
            continue
        if not companies:
            invalid_deep_files.append(path.name)
        for item in companies:
            try:
                researched[int(item["id"])] = item
            except (KeyError, TypeError, ValueError):
                continue

    rows: list[dict[str, Any]] = []
    for company_id in sorted(source_by_id):
        source = source_by_id[company_id]
        research = researched.get(company_id)
        profile_path = ""
        if research:
            grade = research.get("gate2_grade", "D")
            decision = "开发" if research.get("target_customer") else "不主动开发"
            reason = research.get("decision_reason") or research.get("gate1_reason") or "证据不足"
            if research.get("target_customer") and grade in {"A", "B"}:
                folder = PROFILE_ROOT / country_folder(research, source)
                folder.mkdir(parents=True, exist_ok=True)
                profile = folder / f"{company_id:03d}-{safe_filename(research.get('company') or source['company'])}.md"
                profile.write_text(make_profile(research, source), encoding="utf-8")
                profile_path = str(profile.relative_to(ROOT))
            row = {
                "序号": company_id,
                "公司": research.get("company") or source["company"],
                "国家/地区": research.get("country") or source.get("location", ""),
                "官网": research.get("official_website") or source.get("official_url", ""),
                "供应链位置": research.get("supply_chain_role", "Unknown"),
                "闸门结果": grade,
                "是否目标客户": "是" if research.get("target_customer") else "否",
                "Fit": research.get("fit_score", ""),
                "Value": research.get("value_score", ""),
                "价值梯队": research.get("value_tier", ""),
                "首选联系人": (research.get("contacts") or [{}])[0].get("name", "待确认"),
                "首选切入角度": (research.get("china_supplier_entry_angles") or [{}])[0].get("angle", ""),
                "取舍决定": decision,
                "取舍原因": reason,
                "建议动作": research.get("recommended_action", ""),
                "档案": profile_path,
            }
        else:
            gate = source.get("gate1") or source.get("status") or "未完成"
            reason = source.get("reason") or "深度调查未完成"
            row = {
                "序号": company_id,
                "公司": source["company"],
                "国家/地区": source.get("location", ""),
                "官网": source.get("official_url", ""),
                "供应链位置": "Unknown",
                "闸门结果": gate,
                "是否目标客户": "否",
                "Fit": "",
                "Value": "",
                "价值梯队": "",
                "首选联系人": "待确认",
                "首选切入角度": "",
                "取舍决定": "淘汰" if gate == "Reject" else "待复核",
                "取舍原因": reason,
                "建议动作": "不投入深调研" if gate == "Reject" else "补查",
                "档案": "",
            }
        rows.append(row)

    counts: dict[str, int] = {}
    for row in rows:
        counts[row["闸门结果"]] = counts.get(row["闸门结果"], 0) + 1
    target_count = sum(1 for row in rows if row["是否目标客户"] == "是")

    summary_lines = [
        "# LinkedIn 品牌商品公司批量调查——全部客户取舍原因",
        "",
        "> 调查日期：2026-08-13  ",
        f"> 总公司数：{len(rows)}；目标客户：{target_count}；非目标/淘汰/待复核：{len(rows) - target_count}  ",
        f"> 分级统计：{'；'.join(f'{key}={value}' for key, value in sorted(counts.items()))}  ",
        "> 判断基准：Wischos 中国定制金属礼品；MOQ 100；FOB US$18–50/套；交期 25–35 天。",
        "",
        "## 全部客户取舍列表",
        "",
        "| # | 公司 | 地区 | 供应链位置 | 分级 | 目标 | Fit | Value | 取舍原因 | 档案 |",
        "|---:|---|---|---|---|---|---:|---:|---|---|",
    ]
    for row in rows:
        profile_link = f"[查看](</mnt/d/Project/Wischos trading website/{row['档案']}>)" if row["档案"] else "—"
        summary_lines.append(
            f"| {row['序号']} | {md(row['公司'])} | {md(row['国家/地区'])} | {md(row['供应链位置'])} | "
            f"{md(row['闸门结果'])} | {row['是否目标客户']} | {md(row['Fit'])} | {md(row['Value'])} | {md(row['取舍原因'])} | {profile_link} |"
        )
    if invalid_deep_files:
        summary_lines.extend(["", "## 需技术复核的深调研批次", "", bullets(invalid_deep_files)])
    SUMMARY_MD.write_text("\n".join(summary_lines) + "\n", encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "全部客户取舍"
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row[header] for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    target_fill = PatternFill("solid", fgColor="E2F0D9")
    reject_fill = PatternFill("solid", fgColor="FCE4D6")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index in range(2, ws.max_row + 1):
        is_target = ws.cell(row_index, headers.index("是否目标客户") + 1).value == "是"
        fill = target_fill if is_target else reject_fill
        for cell in ws[row_index]:
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {
        "序号": 8, "公司": 34, "国家/地区": 26, "官网": 38, "供应链位置": 22,
        "闸门结果": 12, "是否目标客户": 14, "Fit": 9, "Value": 9, "价值梯队": 18,
        "首选联系人": 24, "首选切入角度": 40, "取舍决定": 14, "取舍原因": 70,
        "建议动作": 42, "档案": 55,
    }
    for index, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(header, 20)
    ws.sheet_view.showGridLines = False
    SUMMARY_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SUMMARY_XLSX)

    print(json.dumps({"rows": len(rows), "researched": len(researched), "targets": target_count, "counts": counts, "invalid_deep_files": invalid_deep_files}, ensure_ascii=False))


if __name__ == "__main__":
    main()
