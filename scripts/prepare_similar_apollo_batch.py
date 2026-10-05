#!/usr/bin/env python3
"""Prepare the Similar/Apollo 2026-09 batch for staged customer research."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "客户系统/02_原始名单与批次/历史筛选与汇总"
    / "similar apollo客户待查_客户交叉比对_2026-09-09.xlsx"
)
MASTER = ROOT / "客户系统/00_主索引/Wischos_客户主索引.xlsx"
OUT = ROOT / ".firecrawl/similar-apollo-2026-09-10"

LEGAL = re.compile(
    r"\b(?:pty\.?\s*ltd\.?|ltd\.?|limited|inc\.?|incorporated|llc|corp\.?|"
    r"corporation|co\.?|company|group|gmbh|srl|bv|a/s|ag|e\.?k\.?)\b",
    re.I,
)
NOISE = re.compile(r"[^a-z0-9]+")

COUNTRY_FOLDER = {
    "united kingdom": "英国",
    "australia": "澳大利亚",
    "new zealand": "新西兰",
    "ireland": "爱尔兰",
    "germany": "欧洲",
    "switzerland": "欧洲",
    "sweden": "欧洲",
    "united arab emirates": "UAE",
    "spain": "欧洲",
    "italy": "欧洲",
    "austria": "欧洲",
    "denmark": "欧洲",
    "netherlands": "欧洲",
    "belgium": "欧洲",
    "france": "欧洲",
    "united states": "美国",
    "us": "美国",
    "pakistan": "巴基斯坦",
    "colombia": "哥伦比亚",
}

OUT_OF_SCOPE = {
    "pakistan",
    "united states",
    "us",
    "colombia",
}

# Apollo sometimes appends a descriptor or market suffix to an existing
# account name.  These pairs were verified by the same official domain and
# market; use the established profile instead of creating a parallel folder.
FORMAL_NAME_ALIASES = {
    "Inck Merchandise - The Branded Merchandise Specialists": "Inck Merchandise",
    "TurnKey Promotions - Australia": "TurnKey Promotions",
}

# Snapshot of the records verified as pre-existing before this batch writes any
# profiles.  Keeping this explicit makes reruns idempotent: newly created batch
# profiles must not cause every A/B row to be reclassified as "Existing".
VERIFIED_EXISTING_IDS = {
    3, 5, 6, 7, 8, 11, 12, 14, 16, 19, 20, 21, 25, 26, 27, 31, 39,
    41, 53, 60, 61, 65, 74, 87, 91, 92,
}


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def normal_name(value: str) -> str:
    value = re.sub(r"\([^)]*\)", " ", value.lower())
    value = value.replace("®", " ").replace("•", " ")
    value = LEGAL.sub(" ", value)
    return NOISE.sub("", value)


def country_from_location(location: str) -> str:
    lower = location.lower().strip()
    aliases = {
        "california": "united states",
        "michigan": "united states",
        "new jersey": "united states",
        "north carolina": "united states",
        "texas": "united states",
        "wisconsin": "united states",
        "delaware": "united states",
        "wales": "united kingdom",
        "england": "united kingdom",
    }
    for token, country in aliases.items():
        if token in lower:
            return country
    for country in COUNTRY_FOLDER:
        if country in lower:
            return "united states" if country == "us" else country
    # Rows such as "Ireland" contain no comma.
    return lower


def formal_country(path: str) -> str:
    parts = Path(path.splitlines()[0]).parts if path else ()
    try:
        index = parts.index("01_正式客户档案")
        return parts[index + 1]
    except (ValueError, IndexError):
        return ""


def compatible_formal(input_country: str, path: str, company: str) -> bool:
    if not path:
        return False
    folder = formal_country(path)
    expected = COUNTRY_FOLDER.get(input_country, "")
    if expected and folder == expected:
        return True
    # This profile is the same UK business but was historically filed under 其他市场.
    if normal_name(company) == normal_name("JSM Brand Exposure") and folder == "其他市场":
        return True
    return False


def main() -> None:
    wb = load_workbook(SOURCE, data_only=True, read_only=True)
    source_rows = list(wb["对比结果"].iter_rows(min_row=2, values_only=True))

    master_wb = load_workbook(MASTER, data_only=True, read_only=True)
    ws = master_wb["客户主索引"]
    headers = [clean(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    index = {header: pos for pos, header in enumerate(headers)}
    master_by_key: dict[str, list[tuple]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        key = normal_name(clean(row[index["标准公司名"]]))
        master_by_key.setdefault(key, []).append(row)

    manifest = []
    for row_id, row in enumerate(source_rows, 1):
        name, location, employees, original_status, match_score, original_sources = map(clean, row)
        country = country_from_location(location)
        # Apollo attached a Delaware location to this German e.K. business.  The
        # verified company identity and legal form control the filing country.
        if name == "Proline Werbeartikel e.K.":
            country = "germany"
        lookup_name = FORMAL_NAME_ALIASES.get(name, name)
        matches = master_by_key.get(normal_name(lookup_name), [])
        if row_id not in VERIFIED_EXISTING_IDS:
            matches = []
        formal_matches = []
        for match in matches:
            formal_path = clean(match[index["正式档案路径"]])
            if compatible_formal(country, formal_path, name):
                formal_matches.append(
                    {
                        "customer_id": clean(match[index["客户ID"]]),
                        "standard_name": clean(match[index["标准公司名"]]),
                        "formal_profile": formal_path,
                        "research_status": clean(match[index["调查状态"]]),
                    }
                )

        if formal_matches:
            batch_status = "ExistingFormal"
            batch_reason = "当前客户系统中已有同市场正式档案，保留原档，不重复调查。"
        # Proline is a German-legal-name lead whose Apollo location says Delaware;
        # keep it in identity verification instead of excluding it on that field alone.
        elif country in OUT_OF_SCOPE and name != "Proline Werbeartikel e.K.":
            batch_status = "ScopeExcluded"
            batch_reason = "不在 Wischos 当前目标市场（Australia/EU/UK/Canada/UAE）内，本轮不调查。"
        else:
            batch_status = "Gate1Pending"
            batch_reason = "无同市场正式档案，进入 Gate 1。"

        manifest.append(
            {
                "id": row_id,
                "company": name,
                "location": location,
                "country": country,
                "country_folder": COUNTRY_FOLDER.get(country, "其他市场"),
                "employees_input": employees,
                "original_crosscheck": original_status,
                "original_match_score": match_score,
                "original_sources": original_sources,
                "batch_status": batch_status,
                "batch_reason": batch_reason,
                "formal_matches": formal_matches,
            }
        )

    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    targets = [row for row in manifest if row["batch_status"] == "Gate1Pending"]
    (OUT / "gate1-targets.json").write_text(
        json.dumps(targets, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    counts: dict[str, int] = {}
    for row in manifest:
        counts[row["batch_status"]] = counts.get(row["batch_status"], 0) + 1
    print(json.dumps({"total": len(manifest), **counts}, ensure_ascii=False))
    for row in manifest:
        print(f"{row['id']:02d}\t{row['batch_status']}\t{row['company']}\t{row['location']}")


if __name__ == "__main__":
    main()
