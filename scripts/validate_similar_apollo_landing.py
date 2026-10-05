#!/usr/bin/env python3
"""Validate the landed Similar/Apollo batch and its customer-system invariants."""

from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BATCH = ROOT / ".firecrawl/similar-apollo-2026-09-10"
SOURCE = ROOT / "客户系统/02_原始名单与批次/历史筛选与汇总/similar apollo客户待查_客户交叉比对_2026-09-09.xlsx"
RESULT = ROOT / "客户系统/02_原始名单与批次/历史筛选与汇总/similar apollo客户待查_客户交叉比对_2026-09-09_调查结果.xlsx"
MASTER = ROOT / "客户系统/00_主索引/Wischos_客户主索引.xlsx"
CUSTOMS = ["ImportYeti", "Customs.report", "52wmb", "Volza", "ImportGenius", "Panjiva", "Datamyne", "TradeSNS"]


def rows(file: Path, sheet: str) -> list[tuple]:
    ws = load_workbook(file, read_only=True, data_only=True)[sheet]
    return list(ws.iter_rows(min_row=2, values_only=True))


def main() -> None:
    errors: list[str] = []
    source_rows = rows(SOURCE, "对比结果")
    result_rows = rows(RESULT, "调查结果")
    if len(source_rows) != 95 or len(result_rows) != len(source_rows):
        errors.append(f"结果行数错误 source={len(source_rows)} result={len(result_rows)}")

    counts = Counter(str(row[6] or "").strip() for row in result_rows)
    gate3 = json.loads((BATCH / "gate3-results.json").read_text(encoding="utf-8"))["companies"]
    manifest = {int(row["id"]): row for row in json.loads((BATCH / "manifest.json").read_text(encoding="utf-8"))}
    final_ab = {
        int(company["id"]): company for company in gate3
        if company["gate2_grade"] in {"A", "B"}
        and manifest[int(company["id"])]["batch_status"] != "ExistingFormal"
    }
    landed_ab = {int(row[0]): row for row in result_rows if str(row[6] or "").strip() in {"A", "B"}}
    if set(landed_ab) != set(final_ab):
        errors.append(f"A/B落档ID不一致 expected={sorted(final_ab)} actual={sorted(landed_ab)}")

    for company_id, row in landed_ab.items():
        path = Path(str(row[15]))
        if not path.exists():
            errors.append(f"id={company_id} 正式档案不存在：{path}")
            continue
        text = path.read_text(encoding="utf-8")
        for label in CUSTOMS:
            if f"| {label} |" not in text:
                errors.append(f"id={company_id} 缺少海关平台 {label}")
        match = re.search(r"\*\*字符数:\s*(\d+)\*\*", text)
        if not match or int(match.group(1)) > 300:
            errors.append(f"id={company_id} LinkedIn邀请字符数无效")
        if "17/17 项完成；0 项失败" not in text:
            errors.append(f"id={company_id} 外部证据摘要不完整")

    for company in gate3:
        if len(company.get("china_supplier_entry_angles") or []) != 3:
            errors.append(f"id={company['id']} 切入角度不是3个")
        if not company.get("research_sources"):
            errors.append(f"id={company['id']} 没有研究来源")

    master_rows = rows(MASTER, "客户主索引")
    ids = [str(row[0] or "").strip() for row in master_rows]
    if len(ids) != len(set(ids)):
        errors.append(f"主索引客户ID不唯一 rows={len(ids)} unique={len(set(ids))}")

    if errors:
        raise SystemExit("\n".join(errors))
    print(json.dumps({
        "source_rows": len(source_rows),
        "result_rows": len(result_rows),
        "result_counts": dict(counts),
        "formal_profiles": len(final_ab),
        "master_rows": len(master_rows),
        "master_unique_ids": len(set(ids)),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
